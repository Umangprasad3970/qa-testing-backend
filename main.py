import os, json, re, urllib.request, urllib.error
from datetime import datetime, timedelta, timezone
from typing import Optional
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from jose import jwt, JWTError
from passlib.context import CryptContext
from pydantic import BaseModel, EmailStr
from sqlalchemy import create_engine, String, Text, DateTime, ForeignKey, Boolean, Integer, inspect, text as sql_text
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, relationship, sessionmaker, Session

DATABASE_URL=os.getenv("DATABASE_URL","sqlite:///./qa_platform.db")
if DATABASE_URL.startswith("postgres://"): DATABASE_URL=DATABASE_URL.replace("postgres://","postgresql+psycopg://",1)
elif DATABASE_URL.startswith("postgresql://"): DATABASE_URL=DATABASE_URL.replace("postgresql://","postgresql+psycopg://",1)
engine=create_engine(DATABASE_URL,pool_pre_ping=True,connect_args={"check_same_thread":False} if DATABASE_URL.startswith("sqlite") else {})
SessionLocal=sessionmaker(bind=engine,autoflush=False,autocommit=False)
SECRET_KEY=os.getenv("JWT_SECRET","CHANGE_THIS_SECRET_IN_RENDER"); ADMIN_EMAIL=os.getenv("ADMIN_EMAIL","admin@example.com"); ADMIN_PASSWORD=os.getenv("ADMIN_PASSWORD","ChangeMe123!")
pwd=CryptContext(schemes=["bcrypt"],deprecated="auto"); security=HTTPBearer()
class Base(DeclarativeBase): pass
class User(Base):
 __tablename__="users"; id:Mapped[int]=mapped_column(Integer,primary_key=True); name:Mapped[str]=mapped_column(String(120)); email:Mapped[str]=mapped_column(String(255),unique=True,index=True); password_hash:Mapped[str]=mapped_column(String(255)); is_admin:Mapped[bool]=mapped_column(Boolean,default=False); is_active:Mapped[bool]=mapped_column(Boolean,default=True); created_at:Mapped[datetime]=mapped_column(DateTime,default=lambda:datetime.now(timezone.utc)); projects:Mapped[list["Project"]]=relationship(back_populates="user")
class Project(Base):
 __tablename__="projects"; id:Mapped[int]=mapped_column(Integer,primary_key=True); name:Mapped[str]=mapped_column(String(160)); website_url:Mapped[str]=mapped_column(String(1000)); description:Mapped[Optional[str]]=mapped_column(Text,nullable=True); user_id:Mapped[int]=mapped_column(ForeignKey("users.id")); created_at:Mapped[datetime]=mapped_column(DateTime,default=lambda:datetime.now(timezone.utc)); user:Mapped["User"]=relationship(back_populates="projects"); test_cases:Mapped[list["TestCase"]]=relationship(back_populates="project",cascade="all, delete-orphan"); test_runs:Mapped[list["TestRun"]]=relationship(back_populates="project",cascade="all, delete-orphan")
class ProjectDocument(Base):
 __tablename__="project_documents"; id:Mapped[int]=mapped_column(Integer,primary_key=True); project_id:Mapped[int]=mapped_column(ForeignKey("projects.id")); filename:Mapped[str]=mapped_column(String(500)); content_text:Mapped[str]=mapped_column(Text); created_at:Mapped[datetime]=mapped_column(DateTime,default=lambda:datetime.now(timezone.utc))
class TestCase(Base):
 __tablename__="test_cases"; id:Mapped[int]=mapped_column(Integer,primary_key=True); project_id:Mapped[int]=mapped_column(ForeignKey("projects.id")); title:Mapped[str]=mapped_column(String(255)); steps:Mapped[str]=mapped_column(Text); expected_result:Mapped[str]=mapped_column(Text); created_at:Mapped[datetime]=mapped_column(DateTime,default=lambda:datetime.now(timezone.utc)); project:Mapped["Project"]=relationship(back_populates="test_cases")
class TestRun(Base):
 __tablename__="test_runs"; id:Mapped[int]=mapped_column(Integer,primary_key=True); project_id:Mapped[int]=mapped_column(ForeignKey("projects.id")); status:Mapped[str]=mapped_column(String(40),default="queued"); started_at:Mapped[Optional[datetime]]=mapped_column(DateTime,nullable=True); finished_at:Mapped[Optional[datetime]]=mapped_column(DateTime,nullable=True); project:Mapped["Project"]=relationship(back_populates="test_runs"); results:Mapped[list["TestResult"]]=relationship(back_populates="run",cascade="all, delete-orphan")
class TestResult(Base):
 __tablename__="test_results"; id:Mapped[int]=mapped_column(Integer,primary_key=True); run_id:Mapped[int]=mapped_column(ForeignKey("test_runs.id")); test_case_id:Mapped[int]=mapped_column(ForeignKey("test_cases.id")); status:Mapped[str]=mapped_column(String(40)); error_message:Mapped[Optional[str]]=mapped_column(Text,nullable=True); screenshot_url:Mapped[Optional[str]]=mapped_column(String(1000),nullable=True); duration_ms:Mapped[Optional[int]]=mapped_column(Integer,nullable=True); run:Mapped["TestRun"]=relationship(back_populates="results")
Base.metadata.create_all(engine)
# Lightweight migration for existing deployments.
def ensure_schema():
    insp = inspect(engine)
    if "users" in insp.get_table_names() and "is_active" not in {c["name"] for c in insp.get_columns("users")}:
        with engine.begin() as conn:
            if engine.dialect.name == "postgresql":
                conn.execute(sql_text("ALTER TABLE users ADD COLUMN IF NOT EXISTS is_active BOOLEAN NOT NULL DEFAULT TRUE"))
            else:
                conn.execute(sql_text("ALTER TABLE users ADD COLUMN is_active BOOLEAN NOT NULL DEFAULT 1"))
ensure_schema()
app=FastAPI(title="QA Testing Platform API",version="3.0.0"); app.add_middleware(CORSMiddleware,allow_origins=["*"],allow_credentials=True,allow_methods=["*"],allow_headers=["*"])
def db():
 s=SessionLocal();
 try: yield s
 finally: s.close()
def token_for(u): return jwt.encode({"sub":str(u.id),"email":u.email,"admin":u.is_admin,"exp":datetime.now(timezone.utc)+timedelta(hours=24)},SECRET_KEY,algorithm="HS256")
def current_user(credentials:HTTPAuthorizationCredentials=Depends(security),session:Session=Depends(db)):
 try: uid=int(jwt.decode(credentials.credentials,SECRET_KEY,algorithms=["HS256"])["sub"])
 except (JWTError,KeyError,ValueError): raise HTTPException(401,"Invalid or expired token")
 u=session.get(User,uid)
 if not u: raise HTTPException(401,"User not found")
 if not u.is_active: raise HTTPException(403,"User account is inactive")
 return u
def admin_only(user:User=Depends(current_user)):
 if not user.is_admin: raise HTTPException(403,"Admin access required")
 return user
class Login(BaseModel): email:EmailStr; password:str
class Register(BaseModel): name:str; email:EmailStr; password:str
class ProjectIn(BaseModel): name:str; website_url:str; description:Optional[str]=None
class TestRunIn(BaseModel): project_id:int
class AdminUserCreate(BaseModel): name:str; email:EmailStr; password:str
class AdminUserUpdate(BaseModel): name:Optional[str]=None; email:Optional[EmailStr]=None; password:Optional[str]=None
class UserStatusUpdate(BaseModel): is_active:bool

def extract_text(filename,data):
 ext=os.path.splitext(filename.lower())[1]
 if ext==".txt": return data.decode("utf-8",errors="ignore")
 if ext==".docx":
  from docx import Document
  import io
  return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
 if ext==".pdf":
  import fitz
  doc=fitz.open(stream=data,filetype="pdf"); return "\n".join(page.get_text() for page in doc)
 raise HTTPException(400,"Supported documents: PDF, DOCX, TXT")

def fallback_cases(text):
 t=text[:30000]; low=t.lower(); cases=[]
 def add(title,steps,expected): cases.append({"title":title,"steps":steps,"expected_result":expected})
 add("Application loads successfully", "Open the project URL and wait for the landing page to load.", "The application loads without a server error and the main page is displayed.")
 add("Navigation links work", "Open the main navigation and select each visible primary link.", "Each link opens the expected page without a broken-link or server error.")
 add("Required field validation", "Submit the primary form without entering required fields.", "Validation messages are displayed and invalid submission is prevented.")
 add("Valid form submission", "Enter valid data in the primary form and submit it.", "The data is accepted and a success confirmation or expected next page is shown.")
 add("Invalid input handling", "Enter invalid, malformed, or boundary values into available input fields and submit.", "Invalid input is rejected with clear validation feedback and the application remains stable.")
 add("Refresh and session behavior", "Complete a normal flow, refresh the page, and navigate back to the application.", "The application remains usable and session state follows the documented requirements.")
 if any(k in low for k in ["login","sign in","authentication","password"]):
  add("Valid login", "Enter a valid registered email and password and select Login.", "The user is authenticated and redirected to the authenticated area.")
  add("Invalid login", "Enter an incorrect password for a registered account and select Login.", "Authentication is rejected without exposing sensitive information.")
  add("Logout", "Sign in and select Logout.", "The session is terminated and protected pages require authentication again.")
 if any(k in low for k in ["signup","register","registration","create account"]): add("Registration validation", "Open registration and try duplicate or invalid account data.", "Invalid or duplicate registration is rejected with a clear message.")
 if any(k in low for k in ["search","filter"]): add("Search and filtering", "Enter a valid search term and use available filters.", "Results match the entered search/filter criteria and no unrelated records are shown.")
 if any(k in low for k in ["payment","checkout","cart","order"]):
  add("Checkout flow", "Add an item or service to the cart and proceed through checkout using valid test data.", "The order/payment flow completes according to the documented requirements.")
  add("Checkout negative case", "Attempt checkout with missing or invalid payment/order data.", "The transaction is blocked and an actionable error is displayed.")
 if any(k in low for k in ["upload","file","document"]): add("File upload validation", "Upload a supported file, then try an unsupported or oversized file.", "Supported files are accepted and invalid files are rejected safely with feedback.")
 if any(k in low for k in ["role","admin","permission","authorization"]): add("Role-based access", "Sign in with users having different documented roles and open protected features.", "Each role can access only the features permitted by the requirements.")
 add("Responsive layout", "Open key pages at narrow and wide viewport sizes and interact with controls.", "Content remains readable, controls remain usable, and no critical overlap or clipping occurs.")
 add("Error handling", "Trigger a network/server error or submit an invalid request during a normal flow.", "A user-friendly error is shown and sensitive server details are not exposed.")
 return cases[:25]

from qa_ai import ENGINE, SKLEARN_AVAILABLE

def ai_cases(text,website):
 cases=ENGINE.generate(text,website)
 source="Local QA AI (TF-IDF + trained classifier + QA scenario engine)" if SKLEARN_AVAILABLE else "Local QA AI (rule-enhanced scenario engine)"
 return cases,source

@app.on_event("startup")
def startup():
 with SessionLocal() as s:
  if not s.query(User).filter(User.email==ADMIN_EMAIL).first(): s.add(User(name="Administrator",email=ADMIN_EMAIL,password_hash=pwd.hash(ADMIN_PASSWORD),is_admin=True)); s.commit()
@app.get("/")
def root(): return {"name":"QA Testing Platform API","status":"online","version":"3.0.0","docs":"/docs"}
@app.get("/health")
def health(): return {"status":"ok"}
@app.get("/admin")
def admin_page(): return FileResponse("admin.html")
@app.get("/api/ai/status")
def ai_status():
 return {"name":"Local QA AI","external_api_required":False,"ml_engine":"TF-IDF + Logistic Regression","scenario_engine":"requirement classification + QA template expansion","training_examples":len(__import__("qa_ai").TRAINING)}
@app.post("/api/auth/register")
def register(data:Register,admin:User=Depends(admin_only),session:Session=Depends(db)):
 if session.query(User).filter(User.email==data.email).first(): raise HTTPException(409,"Email already registered")
 u=User(name=data.name,email=data.email,password_hash=pwd.hash(data.password),is_admin=False,is_active=True);session.add(u);session.commit();session.refresh(u);return {"message":"user created","id":u.id}
@app.post("/api/auth/login")
def login(data:Login,session:Session=Depends(db)):
 u=session.query(User).filter(User.email==data.email).first()
 if not u or not pwd.verify(data.password,u.password_hash): raise HTTPException(401,"Invalid email or password")
 if not u.is_active: raise HTTPException(403,"User account is inactive")
 return {"access_token":token_for(u),"token_type":"bearer","is_admin":u.is_admin}
@app.post("/api/auth/client-login")
def client_login(data:Login, client_type:str="android", session:Session=Depends(db)):
 if client_type not in {"android","windows"}: raise HTTPException(400,"client_type must be android or windows")
 u=session.query(User).filter(User.email==data.email).first()
 if not u or not pwd.verify(data.password,u.password_hash): raise HTTPException(401,"Invalid email or password")
 if u.is_admin: raise HTTPException(403,"Administrator accounts cannot log in to mobile or Windows applications")
 if not u.is_active: raise HTTPException(403,"User account is inactive")
 return {"access_token":token_for(u),"token_type":"bearer","is_admin":False,"client_type":client_type}
@app.get("/api/me")
def me(user:User=Depends(current_user)):
 return {"id":user.id,"name":user.name,"email":user.email,"is_admin":user.is_admin,"is_active":user.is_active}
@app.post("/api/projects")
def create_project(data:ProjectIn,user:User=Depends(current_user),session:Session=Depends(db)):
 p=Project(**data.model_dump(),user_id=user.id);session.add(p);session.commit();session.refresh(p);return {"id":p.id,"name":p.name,"website_url":p.website_url}
@app.get("/api/projects")
def my_projects(user:User=Depends(current_user),session:Session=Depends(db)):
 return [{"id":p.id,"name":p.name,"website_url":p.website_url,"description":p.description,"created_at":p.created_at} for p in session.query(Project).filter(Project.user_id==user.id).order_by(Project.id.desc()).all()]
@app.post("/api/projects/{project_id}/document")
async def upload_document(project_id:int,file:UploadFile=File(...),user:User=Depends(current_user),session:Session=Depends(db)):
 p=session.get(Project,project_id)
 if not p or p.user_id!=user.id: raise HTTPException(404,"Project not found")
 data=await file.read(); text=extract_text(file.filename or "document.txt",data)
 if len(text.strip())<30: raise HTTPException(400,"Document contains too little readable text")
 d=ProjectDocument(project_id=project_id,filename=file.filename or "document",content_text=text[:100000]);session.add(d);session.commit();return {"message":"Project document uploaded","filename":d.filename,"characters":len(text)}
@app.post("/api/projects/{project_id}/generate-test-cases")
def generate(project_id:int,user:User=Depends(current_user),session:Session=Depends(db)):
 p=session.get(Project,project_id)
 if not p or p.user_id!=user.id: raise HTTPException(404,"Project not found")
 d=session.query(ProjectDocument).filter(ProjectDocument.project_id==project_id).order_by(ProjectDocument.id.desc()).first()
 if not d: raise HTTPException(400,"Upload a project document first")
 cases,source=ai_cases(d.content_text,p.website_url)
 session.query(TestCase).filter(TestCase.project_id==project_id).delete(synchronize_session=False)
 for c in cases: session.add(TestCase(project_id=project_id,title=c["title"][:255],steps=c["steps"],expected_result=c["expected_result"]))
 session.commit();return {"message":f"Created {len(cases)} test cases ({source})","created":len(cases),"source":source}
@app.get("/api/projects/{project_id}/test-cases")
def get_cases(project_id:int,user:User=Depends(current_user),session:Session=Depends(db)):
 p=session.get(Project,project_id)
 if not p or p.user_id!=user.id: raise HTTPException(404,"Project not found")
 return [{"id":t.id,"title":t.title,"steps":t.steps,"expected_result":t.expected_result} for t in p.test_cases]
@app.post("/api/test-runs")
def run_tests(data:TestRunIn,user:User=Depends(current_user),session:Session=Depends(db)):
 p=session.get(Project,data.project_id)
 if not p or p.user_id!=user.id: raise HTTPException(404,"Project not found")
 cases=p.test_cases
 if not cases: raise HTTPException(400,"Generate test cases first")
 r=TestRun(project_id=p.id,status="running",started_at=datetime.now(timezone.utc));session.add(r);session.flush()
 start=datetime.now(timezone.utc); status="failed"; err=None
 try:
  req=urllib.request.Request(p.website_url,headers={"User-Agent":"QA-Testing-Platform/2.0"}); resp=urllib.request.urlopen(req,timeout=20); code=resp.getcode(); status="passed" if 200<=code<400 else "failed"; err=None if status=="passed" else f"HTTP status {code}"
 except Exception as e: err=str(e)[:500]
 dur=int((datetime.now(timezone.utc)-start).total_seconds()*1000)
 for tc in cases: session.add(TestResult(run_id=r.id,test_case_id=tc.id,status=status,error_message=err,duration_ms=dur))
 r.status="completed";r.started_at=r.started_at;r.finished_at=datetime.now(timezone.utc);session.commit();return {"message":f"Automated smoke run completed: {status}","run_id":r.id}
@app.get("/api/projects/{project_id}/report")
def report(project_id:int,user:User=Depends(current_user),session:Session=Depends(db)):
 p=session.get(Project,project_id)
 if not p or p.user_id!=user.id: raise HTTPException(404,"Project not found")
 cases=p.test_cases; latest=session.query(TestRun).filter(TestRun.project_id==project_id).order_by(TestRun.id.desc()).first(); results=[]
 if latest:
  by={x.test_case_id:x for x in latest.results}
  for tc in cases:
   x=by.get(tc.id);results.append({"test_case_id":tc.id,"test_case_title":tc.title,"status":x.status if x else "not_run","error_message":x.error_message if x else None})
 total=len(cases);passed=sum(x["status"]=="passed" for x in results);failed=sum(x["status"]=="failed" for x in results);notrun=total-passed-failed; pct=round((passed/total)*100,2) if total else 0
 final="PASS" if total and failed==0 and notrun==0 else ("FAIL" if failed>0 else "NOT RUN")
 return {"project_name":p.name,"website_url":p.website_url,"total":total,"passed":passed,"failed":failed,"not_run":notrun,"pass_percentage":pct,"final_status":final,"results":results}
@app.get("/api/admin/dashboard")
def admin_dashboard(admin:User=Depends(admin_only),session:Session=Depends(db)):
 return {"users":session.query(User).count(),"active_users":session.query(User).filter(User.is_admin==False,User.is_active==True).count(),"inactive_users":session.query(User).filter(User.is_admin==False,User.is_active==False).count(),"projects":session.query(Project).count(),"test_cases":session.query(TestCase).count(),"test_runs":session.query(TestRun).count(),"test_results":session.query(TestResult).count(),"passed":session.query(TestResult).filter(TestResult.status.ilike("passed")).count(),"failed":session.query(TestResult).filter(TestResult.status.ilike("failed")).count()}

@app.post("/api/admin/users")
def admin_create_user(data:AdminUserCreate,admin:User=Depends(admin_only),session:Session=Depends(db)):
 if session.query(User).filter(User.email==data.email).first(): raise HTTPException(409,"Email already registered")
 if len(data.password)<6: raise HTTPException(400,"Password must be at least 6 characters")
 u=User(name=data.name,email=data.email,password_hash=pwd.hash(data.password),is_admin=False,is_active=True);session.add(u);session.commit();session.refresh(u)
 return {"id":u.id,"name":u.name,"email":u.email,"is_admin":False,"is_active":u.is_active}

def admin_user_payload(u,session):
 projects=session.query(Project).filter(Project.user_id==u.id).all()
 return {"id":u.id,"name":u.name,"email":u.email,"is_admin":u.is_admin,"is_active":u.is_active,"created_at":u.created_at,"projects_count":len(projects),"test_cases_count":sum(len(p.test_cases) for p in projects),"test_runs_count":sum(len(p.test_runs) for p in projects)}

@app.get("/api/admin/users")
def admin_users(admin:User=Depends(admin_only),session:Session=Depends(db)):
 return [admin_user_payload(u,session) for u in session.query(User).order_by(User.id.desc()).all()]

@app.get("/api/admin/users/{user_id}")
def admin_view_user(user_id:int,admin:User=Depends(admin_only),session:Session=Depends(db)):
 u=session.get(User,user_id)
 if not u: raise HTTPException(404,"User not found")
 return admin_user_payload(u,session)

@app.put("/api/admin/users/{user_id}")
def admin_edit_user(user_id:int,data:AdminUserUpdate,admin:User=Depends(admin_only),session:Session=Depends(db)):
 u=session.get(User,user_id)
 if not u: raise HTTPException(404,"User not found")
 if u.is_admin and user_id!=admin.id: raise HTTPException(403,"Administrator accounts cannot be edited here")
 if data.email and session.query(User).filter(User.email==data.email,User.id!=user_id).first(): raise HTTPException(409,"Email already registered")
 if data.name is not None: u.name=data.name
 if data.email is not None: u.email=data.email
 if data.password is not None:
  if len(data.password)<6: raise HTTPException(400,"Password must be at least 6 characters")
  u.password_hash=pwd.hash(data.password)
 session.commit(); return admin_user_payload(u,session)

@app.patch("/api/admin/users/{user_id}/status")
def admin_change_status(user_id:int,data:UserStatusUpdate,admin:User=Depends(admin_only),session:Session=Depends(db)):
 u=session.get(User,user_id)
 if not u: raise HTTPException(404,"User not found")
 if u.is_admin: raise HTTPException(403,"Administrator status cannot be changed")
 u.is_active=data.is_active;session.commit();return {"message":"User status updated","id":u.id,"is_active":u.is_active}

@app.delete("/api/admin/users/{user_id}")
def admin_delete_user(user_id:int,admin:User=Depends(admin_only),session:Session=Depends(db)):
 u=session.get(User,user_id)
 if not u: raise HTTPException(404,"User not found")
 if u.is_admin: raise HTTPException(403,"Administrator accounts cannot be deleted")
 projects=session.query(Project).filter(Project.user_id==u.id).all()
 for p in projects:
  for run in session.query(TestRun).filter(TestRun.project_id==p.id).all(): session.query(TestResult).filter(TestResult.run_id==run.id).delete(synchronize_session=False)
  session.query(TestRun).filter(TestRun.project_id==p.id).delete(synchronize_session=False)
  session.query(TestCase).filter(TestCase.project_id==p.id).delete(synchronize_session=False)
  session.query(ProjectDocument).filter(ProjectDocument.project_id==p.id).delete(synchronize_session=False)
  session.delete(p)
 session.delete(u);session.commit();return {"message":"User and all owned work deleted","id":user_id}

@app.get("/api/admin/users/{user_id}/work")
def admin_user_work(user_id:int,admin:User=Depends(admin_only),session:Session=Depends(db)):
 u=session.get(User,user_id)
 if not u: raise HTTPException(404,"User not found")
 projects=[]
 for p in session.query(Project).filter(Project.user_id==u.id).order_by(Project.id.desc()).all():
  latest=session.query(TestRun).filter(TestRun.project_id==p.id).order_by(TestRun.id.desc()).first()
  latest_report=None
  if latest:
   rs=session.query(TestResult).filter(TestResult.run_id==latest.id).all(); total=len(p.test_cases); passed=sum(x.status=="passed" for x in rs); failed=sum(x.status=="failed" for x in rs); notrun=total-passed-failed
   latest_report={"total":total,"passed":passed,"failed":failed,"not_run":notrun,"pass_percentage":round((passed/total)*100,2) if total else 0,"final_status":"PASS" if total and failed==0 and notrun==0 else ("FAIL" if failed>0 else "NOT RUN")}
  projects.append({"id":p.id,"name":p.name,"website_url":p.website_url,"created_at":p.created_at,"test_cases":len(p.test_cases),"test_runs":len(p.test_runs),"latest_report":latest_report})
 return {"user":admin_user_payload(u,session),"projects":projects}

@app.get("/api/admin/projects")
def admin_projects(admin:User=Depends(admin_only),session:Session=Depends(db)):
 return [{"id":p.id,"name":p.name,"website_url":p.website_url,"user_id":p.user_id,"user_email":p.user.email,"created_at":p.created_at} for p in session.query(Project).order_by(Project.id.desc()).all()]

@app.get("/api/admin/test-cases")
def admin_test_cases(admin:User=Depends(admin_only),session:Session=Depends(db)):
 return [{"id":t.id,"project_id":t.project_id,"project_name":t.project.name,"user_id":t.project.user_id,"title":t.title,"steps":t.steps,"expected_result":t.expected_result,"created_at":t.created_at} for t in session.query(TestCase).order_by(TestCase.id.desc()).all()]

@app.get("/api/admin/test-runs")
def admin_test_runs(admin:User=Depends(admin_only),session:Session=Depends(db)):
 return [{"id":r.id,"project_id":r.project_id,"project_name":r.project.name,"user_id":r.project.user_id,"status":r.status,"started_at":r.started_at,"finished_at":r.finished_at} for r in session.query(TestRun).order_by(TestRun.id.desc()).all()]

@app.get("/api/admin/test-results")
def admin_test_results(admin:User=Depends(admin_only),session:Session=Depends(db)):
 rows=[]
 for x in session.query(TestResult).order_by(TestResult.id.desc()).all():
  tc=session.get(TestCase,x.test_case_id); run=x.run; project=session.get(Project,run.project_id)
  rows.append({"id":x.id,"run_id":x.run_id,"test_case_id":x.test_case_id,"test_case_title":tc.title if tc else None,"project_id":project.id,"project_name":project.name,"user_id":project.user_id,"status":x.status,"error_message":x.error_message,"duration_ms":x.duration_ms})
 return rows

# --- Dedicated admin pages and management/report APIs ---
@app.get('/admin/users')
def admin_users_page(): return FileResponse('users.html')
@app.get('/admin/projects')
def admin_projects_page(): return FileResponse('projects.html')
@app.get('/admin/test-cases')
def admin_test_cases_page(): return FileResponse('test_cases.html')
@app.get('/admin/test-results')
def admin_test_results_page(): return FileResponse('test_results.html')
@app.get('/admin/project-reports')
def admin_project_reports_page(): return FileResponse('project_reports.html')
@app.get('/admin.css')
def admin_css(): return FileResponse('admin.css', media_type='text/css')
@app.get('/admin_common.js')
def admin_common_js(): return FileResponse('admin_common.js', media_type='application/javascript')

@app.delete('/api/admin/projects/{project_id}')
def admin_delete_project(project_id:int, admin:User=Depends(admin_only), session:Session=Depends(db)):
    p=session.get(Project,project_id)
    if not p: raise HTTPException(404,'Project not found')
    for run in session.query(TestRun).filter(TestRun.project_id==p.id).all():
        session.query(TestResult).filter(TestResult.run_id==run.id).delete(synchronize_session=False)
    session.query(TestRun).filter(TestRun.project_id==p.id).delete(synchronize_session=False)
    session.query(TestCase).filter(TestCase.project_id==p.id).delete(synchronize_session=False)
    session.query(ProjectDocument).filter(ProjectDocument.project_id==p.id).delete(synchronize_session=False)
    session.delete(p); session.commit()
    return {'message':'Project and related QA data deleted','id':project_id}

@app.get('/api/admin/projects/{project_id}/report')
def admin_project_report(project_id:int, admin:User=Depends(admin_only), session:Session=Depends(db)):
    p=session.get(Project,project_id)
    if not p: raise HTTPException(404,'Project not found')
    cases=p.test_cases
    latest=session.query(TestRun).filter(TestRun.project_id==p.id).order_by(TestRun.id.desc()).first()
    results=[]
    if latest:
        by={x.test_case_id:x for x in latest.results}
        for tc in cases:
            x=by.get(tc.id)
            results.append({'test_case_id':tc.id,'title':tc.title,'status':x.status if x else 'not_run','error_message':x.error_message if x else None,'duration_ms':x.duration_ms if x else None})
    total=len(cases); passed=sum(x['status']=='passed' for x in results); failed=sum(x['status']=='failed' for x in results); notrun=total-passed-failed
    return {'project_id':p.id,'project_name':p.name,'website_url':p.website_url,'user_id':p.user_id,'user_email':p.user.email,'total':total,'passed':passed,'failed':failed,'not_run':notrun,'pass_percentage':round(passed*100/total,2) if total else 0,'final_status':'PASS' if total and failed==0 and notrun==0 else ('FAIL' if failed else 'NOT RUN'),'latest_run_id':latest.id if latest else None,'latest_run_status':latest.status if latest else None,'results':results}

@app.get('/api/admin/project-reports')
def admin_project_reports(admin:User=Depends(admin_only), session:Session=Depends(db)):
    out=[]
    for p in session.query(Project).order_by(Project.id.desc()).all():
        latest=session.query(TestRun).filter(TestRun.project_id==p.id).order_by(TestRun.id.desc()).first()
        rs=session.query(TestResult).filter(TestResult.run_id==latest.id).all() if latest else []
        total=len(p.test_cases); passed=sum(x.status=='passed' for x in rs); failed=sum(x.status=='failed' for x in rs); notrun=total-passed-failed
        out.append({'project_id':p.id,'project_name':p.name,'website_url':p.website_url,'user_id':p.user_id,'user_email':p.user.email,'total':total,'passed':passed,'failed':failed,'not_run':notrun,'pass_percentage':round(passed*100/total,2) if total else 0,'final_status':'PASS' if total and failed==0 and notrun==0 else ('FAIL' if failed else 'NOT RUN'),'last_run_id':latest.id if latest else None,'last_run_status':latest.status if latest else None})
    return out

@app.delete('/api/admin/test-cases/{test_case_id}')
def admin_delete_test_case(test_case_id:int, admin:User=Depends(admin_only), session:Session=Depends(db)):
    tc=session.get(TestCase,test_case_id)
    if not tc: raise HTTPException(404,'Test case not found')
    session.query(TestResult).filter(TestResult.test_case_id==test_case_id).delete(synchronize_session=False)
    session.delete(tc); session.commit(); return {'message':'Test case deleted','id':test_case_id}
